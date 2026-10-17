import openpyxl
import os
import json
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
certificates_dir = os.path.join(base_dir, "Certificates")
excel_path = os.path.join(certificates_dir, "TechBETA 2026 2.0 Participation .xlsx")
output_dir = os.path.join(base_dir, "src", "data")
output_json = os.path.join(output_dir, "certificates.json")

folder_map = {
    'UI Verse': 'Ui Verse',
    'GenBuild': 'GenBuild',
    'Logic Trap': 'Logic trap',
    'Idea Forge': 'Idea Forge',
    'Bid & Build': 'Bid & Build',
    'Brand Blitz': 'Brand Blitz'
}

event_types = {
    'UI Verse': 'Technical',
    'GenBuild': 'Technical',
    'Logic Trap': 'Technical',
    'Idea Forge': 'Technical',
    'Bid & Build': 'Non-Technical',
    'Brand Blitz': 'Non-Technical'
}

def build_data():
    if not os.path.exists(excel_path):
        print(f"Error: {excel_path} not found")
        sys.exit(1)
        
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    records = []
    id_counter = 1

    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        folder_name = folder_map.get(sheetname, sheetname)
        folder_path = os.path.join(certificates_dir, folder_name)
        
        for r in range(2, sheet.max_row + 1):
            name = sheet.cell(r, 1).value
            college = sheet.cell(r, 2).value
            event = sheet.cell(r, 3).value or sheetname
            email = sheet.cell(r, 4).value
            
            if not name or not str(name).strip():
                continue
                
            name_str = str(name).strip()
            college_str = str(college).strip() if college else ""
            event_str = str(event).strip() if event else sheetname
            email_str = str(email).strip().lower() if email else ""
            
            filename = f"{name_str}.png"
            full_file_path = os.path.join(folder_path, filename)
            
            if not os.path.exists(full_file_path):
                print(f"Warning: File missing for {name_str} in {folder_name}: {full_file_path}")
            
            file_size = os.path.getsize(full_file_path) if os.path.exists(full_file_path) else 0
            
            records.append({
                "id": f"CERT-{id_counter:04d}",
                "name": name_str,
                "college": college_str,
                "event": event_str,
                "eventType": event_types.get(sheetname, "Technical"),
                "folder": folder_name,
                "email": email_str,
                "filename": filename,
                "fileSize": file_size,
                "verified": True
            })
            id_counter += 1

    os.makedirs(output_dir, exist_ok=True)
    with open(output_json, "w", encoding="utf-8") as f:
        json.dump(records, f, indent=2, ensure_ascii=False)
        
    print(f"Successfully generated {output_json}")
    print(f"Total verified certificates indexed: {len(records)}")
    
    unique_emails = set(r["email"] for r in records if r["email"])
    print(f"Unique recipient emails: {len(unique_emails)}")

if __name__ == '__main__':
    build_data()
